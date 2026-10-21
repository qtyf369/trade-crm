"""
生成300条测试客户数据，用于导入测试
里面故意加一些不合规的数据
"""
import random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime, timedelta

# 基础数据
COUNTRIES = ['美国', '德国', '英国', '法国', '日本', '澳大利亚', '加拿大', '韩国', '印度', '巴西',
             '意大利', '西班牙', '俄罗斯', '墨西哥', '荷兰', '瑞典', '瑞士', '新加坡', '泰国', '马来西亚']

SOURCES = ['阿里巴巴国际站', '展会', 'Google', '其他']

STATUSES = ['新线索', '初步沟通', '已报价未成交', '谈判中', '已成交', '已流失']

LEVELS = ['', 'L0', 'L1', 'L2', 'L3', 'L4', '低质', '优质']

COMPANY_SUFFIXES = ['Trading Co., Ltd.', 'Industrial Co., Ltd.', 'Import & Export Co., Ltd.',
                    'Manufacturing Co., Ltd.', 'Group Co., Ltd.', 'Commerce Co., Ltd.']

FIRST_NAMES = ['John', 'David', 'Michael', 'Sarah', 'Emma', 'James', 'Robert', 'Lisa', 'William', 'Mary',
               'Richard', 'Jennifer', 'Thomas', 'Patricia', 'Charles', 'Linda', 'Daniel', 'Elizabeth',
               'Matthew', 'Susan']

LAST_NAMES = ['Smith', 'Johnson', 'Williams', 'Brown', 'Jones', 'Garcia', 'Miller', 'Davis', 'Rodriguez',
              'Martinez', 'Hernandez', 'Lopez', 'Gonzalez', 'Wilson', 'Anderson', 'Thomas', 'Taylor', 'Moore',
              'Jackson', 'Martin']

PRODUCTS = ['不锈钢餐具', '保温杯', '厨房用具', '五金制品', '不锈钢管材', '餐具套装', '厨具配件',
            '不锈钢板材', '卫浴五金', '门锁五金', '家具五金', '汽车五金配件']

def random_phone():
    return f"+1-{random.randint(200,999)}-{random.randint(200,999)}-{random.randint(1000,9999)}"

def generate_test_data():
    wb = Workbook()
    ws = wb.active
    ws.title = "客户总表"

    # 表头
    headers = [
        "序号", "公司名称", "国家", "来源", "等级", "状态",
        "主要产品", "备注",
        "联系人姓名", "职位", "邮箱", "电话", "社交软件", "社交账号",
        "最近跟进日期", "最近跟进内容",
        "创建时间",
    ]

    # 样式
    header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # 生成300条数据
    row_idx = 2
    for i in range(1, 301):
        # 随机决定这条是不是不合规
        has_bad_country = (i % 17 == 0)  # 17分之一概率国家不对
        has_bad_source = (i % 23 == 0)   # 23分之一概率来源不对
        has_bad_status = (i % 31 == 0)   # 31分之一概率状态不对

        company = f"{random.choice(['ABC', 'XYZ', 'Global', 'Prime', 'Best', 'Top', 'Sun', 'Star', 'Gold', 'Silver'])} {i} {random.choice(COMPANY_SUFFIXES)}"

        country = random.choice(COUNTRIES)
        if has_bad_country:
            # 故意写错国家
            country = random.choice(['米国', '德国汉堡', '美国纽约市', 'England', 'Unknown', ''])

        source = random.choice(SOURCES)
        if has_bad_source:
            # 故意写错来源
            source = random.choice(['Google Ads', 'Facebook', 'LinkedIn', '朋友介绍', '展会名片'])

        status = random.choice(STATUSES)
        if has_bad_status:
            # 故意写错状态
            status = random.choice(['有意向', '已付款', '跟进中', '无效客户', ''])

        level = random.choice(LEVELS)
        product = random.choice(PRODUCTS)
        note = random.choice(['', '老客户', '价格敏感', '量大', '需样品', '长期合作'])

        # 联系人
        contact_name = f"{random.choice(FIRST_NAMES)} {random.choice(LAST_NAMES)}"
        position = random.choice(['采购经理', 'CEO', '采购总监', '产品经理', '老板', ''])
        email = f"{contact_name.lower().replace(' ', '.')}@{random.choice(['gmail.com', 'yahoo.com', 'outlook.com', 'company.com'])}"
        phone = random_phone()
        social = random.choice(['wechat', 'whatsapp', ''])
        social_account = ''
        if social:
            social_account = f"+86 138{random.randint(10000000,99999999)}"

        # 跟进记录
        followup_days_ago = random.randint(0, 30)
        followup_date = (datetime.now() - timedelta(days=followup_days_ago)).strftime('%Y-%m-%d')
        followup_content = random.choice(['发了报价单', '确认样品', '电话沟通', 'WhatsApp联系', '展会见面', '发送产品目录'])

        created_days_ago = random.randint(30, 365)
        created_date = (datetime.now() - timedelta(days=created_days_ago)).strftime('%Y-%m-%d')

        # 写入
        ws.cell(row=row_idx, column=1, value=i)
        ws.cell(row=row_idx, column=2, value=company)
        ws.cell(row=row_idx, column=3, value=country)
        ws.cell(row=row_idx, column=4, value=source)
        ws.cell(row=row_idx, column=5, value=level)
        ws.cell(row=row_idx, column=6, value=status)
        ws.cell(row=row_idx, column=7, value=product)
        ws.cell(row=row_idx, column=8, value=note)
        ws.cell(row=row_idx, column=9, value=contact_name)
        ws.cell(row=row_idx, column=10, value=position)
        ws.cell(row=row_idx, column=11, value=email)
        ws.cell(row=row_idx, column=12, value=phone)
        ws.cell(row=row_idx, column=13, value='微信' if social == 'wechat' else 'WhatsApp' if social == 'whatsapp' else '')
        ws.cell(row=row_idx, column=14, value=social_account)
        ws.cell(row=row_idx, column=15, value=followup_date)
        ws.cell(row=row_idx, column=16, value=followup_content)
        ws.cell(row=row_idx, column=17, value=created_date)

        row_idx += 1

    # 列宽
    col_widths = [6, 28, 12, 16, 8, 14, 22, 28, 12, 12, 28, 16, 12, 16, 13, 35, 12]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    filepath = r"D:\CRM\测试客户数据_300条.xlsx"
    wb.save(filepath)
    print(f"已生成：{filepath}")
    print(f"共300条数据，其中约 {300//17} 条国家不合规，{300//23} 条来源不合规，{300//31} 条状态不合规")

if __name__ == '__main__':
    generate_test_data()
